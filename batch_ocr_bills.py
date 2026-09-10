#!/usr/bin/env python3
"""
批量长截图账单 OCR -> Excel

用法:
    python batch_ocr_bills.py <图片文件夹> [--out 输出.xlsx] [--recursive]

特性:
- 自动跳过像素过小的图（每行文字高度 < 8px 判定为不可识别），并给出报告
- 超长图自动分块 OCR（默认 1600px / 块，重叠 150px）
- 自动按图片宽度推算列边界，并用「金额行」位置校正右列边界
- 跨文件去重：同一笔账单出现在多张重叠截图里只保留一条
- 输出：明细表 + 商户汇总 + 月度汇总 + 异常报告
"""
import argparse, json, os, re, sys, traceback
from collections import Counter, defaultdict

from PIL import Image
from rapidocr_onnxruntime import RapidOCR
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

IMG_EXT = ('.jpg', '.jpeg', '.png', '.webp', '.bmp')
CHUNK_H = 1600
OVERLAP = 150
MIN_TEXT_LINE_H = 8          # 低于该像素高度判定不可识别
REF_WIDTH = 1220             # 参考图宽度，用于按比例推算列边界
REF_LEFT = 190               # 参考图中「头像列」右边界
REF_RIGHT = 930              # 参考图中「金额列」左边界

AMT   = re.compile(r'^[-+]\s?([\d,]+\.\d{2})$')
DATE  = re.compile(r'^(\d{1,2})月(\d{1,2})日\s?(\d{1,2}):(\d{2})$')
PART  = re.compile(r'^(\d{1,2})0?(\d{2}):(\d{2})$')
MONTH = re.compile(r'^(\d{4})年(\d{1,2})月')
SUMM  = re.compile(r'^(收入|支出)\s?[￥¥]\s?([\d,]+\.\d{2})')
FIX   = {'生津英语': '牛津英语'}


# ---------- 1. 图片可用性检查 ----------
def check_readable(path):
    """返回 (是否可用, 说明, 图片对象)"""
    im = Image.open(path).convert('RGB')
    w, h = im.size
    g = im.convert('L')
    # 抽样估算最细文字行高
    sample = max(0, h // 2 - 400)
    crop = g.crop((0, sample, w, min(h, sample + 800)))
    import numpy as np
    a = np.array(crop).astype(float)
    dark = (a < 160).sum(axis=1)
    rows, inrow, start = [], False, 0
    for y, v in enumerate(dark):
        if v > 3 and not inrow:
            inrow, start = True, y
        elif v <= 3 and inrow:
            inrow = False
            rows.append(y - start)
    heights = [r for r in rows if 2 <= r <= 60]
    mh = min(heights) if heights else 999
    if w < 200 or mh < MIN_TEXT_LINE_H:
        return False, f'分辨率不足 {w}x{h}，最细文字行高约 {mh}px（需 ≥{MIN_TEXT_LINE_H}px）', im
    return True, f'{w}x{h}，最细文字行高约 {mh}px', im


# ---------- 2. 分块 OCR ----------
def ocr_image(im, engine):
    w, h = im.size
    boxes = []
    y, idx = 0, 0
    while y < h:
        y1 = min(h, y + CHUNK_H)
        res, _ = engine(__import__('numpy').array(im.crop((0, y, w, y1))))
        if res:
            for box, txt, score in res:
                xs = [p[0] for p in box]
                ys = [p[1] for p in box]
                boxes.append({'y': sum(ys) / 4 + y, 'x0': min(xs), 'x1': max(xs),
                              'text': txt, 'score': float(score)})
        idx += 1
        if y1 >= h:
            break
        y = y1 - OVERLAP
    return boxes


def dedupe_boxes(boxes):
    lines = []
    for b in sorted(boxes, key=lambda r: (r['y'], r['x0'])):
        for L in lines:
            if abs(L['y'] - b['y']) < 18 and abs(L['x0'] - b['x0']) < 40:
                if b['score'] > L['score']:
                    L.update(text=b['text'], score=b['score'], x1=b['x1'])
                break
        else:
            lines.append(dict(b))
    return sorted(lines, key=lambda r: (r['y'], r['x0']))


# ---------- 3. 版面分析 ----------
def detect_columns(lines, width):
    scale = width / REF_WIDTH
    left = REF_LEFT * scale
    right = REF_RIGHT * scale
    amt_x = [L['x0'] for L in lines if AMT.match(L['text'].strip())]
    if len(amt_x) >= 5:
        right = max(min(amt_x) - 0.02 * width, width * 0.5)
    return left, right


# ---------- 4. 组装记录 ----------
def parse_records(lines, left, right, srcname):
    rows = []
    year = month = last_m = last_d = None
    shop = prod = amt = note = None

    def flush(dt=None):
        nonlocal shop, prod, amt, note
        if shop or amt or dt:
            rows.append({'src': srcname, 'dt': dt, 'shop': shop, 'prod': prod,
                         'amt': amt, 'note': note})
        shop = prod = amt = note = None

    for L in lines:
        t = L['text'].strip()
        x0 = L['x0']
        if MONTH.match(t):
            flush()
            m = MONTH.match(t)
            year, month = int(m.group(1)), int(m.group(2))
            continue
        if SUMM.match(t):
            continue
        if DATE.match(t):
            m = DATE.match(t)
            last_m, last_d = int(m.group(1)), int(m.group(2))
            flush(f"{year}-{last_m:02d}-{last_d:02d} {int(m.group(3)):02d}:{m.group(4)}")
            continue
        if PART.match(t):
            m = PART.match(t)
            if last_m:
                flush(f"{year}-{last_m:02d}-{int(m.group(1)):02d} {m.group(2)}:{m.group(3)}")
            continue
        if AMT.match(t):
            amt = '-' + AMT.match(t).group(1)
            continue
        if x0 >= right:          # 右列非金额碎片
            continue
        if x0 < left:            # 头像 / 左侧小字
            continue
        if t.startswith('已优惠'):
            note = t
            continue
        for k, v in FIX.items():
            t = t.replace(k, v)
        if shop is None:
            shop = t
        elif prod is None:
            prod = t
        else:
            prod = prod + ' ' + t
    flush()

    for r in rows:
        r['shop'] = (r['shop'] or '').strip()
        r['prod'] = (r['prod'] or '').strip()
        r['type'] = '退款' if ('退款' in r['shop']) or ('退款' in (r['prod'] or '')) else '支出'
        r['shop_clean'] = r['shop'].replace('退款-', '')
        r['amt_f'] = -abs(float(r['amt'].replace(',', ''))) if r['amt'] else None
        r['flag'] = '缺时间' if not r['dt'] else ('缺金额' if r['amt_f'] is None else '')
    return rows


# ---------- 5. 导出 ----------
HEAD_FILL = PatternFill('solid', fgColor='1F6FEB')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=11)
THIN = Side(style='thin', color='D0D7E2')


def export(all_rows, bad_files, outpath):
    wb = Workbook()
    ws = wb.active
    ws.title = '交易明细'
    cols = ['序号', '交易时间', '类型', '商户名称', '商品说明', '金额(元)', '优惠/备注', '来源文件']
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).fill = HEAD_FILL
        ws.cell(1, c).font = HEAD_FONT
        ws.cell(1, c).alignment = Alignment(horizontal='center', vertical='center')
    for i, r in enumerate(all_rows, 1):
        ws.append([i, r['dt'], r['type'], r['shop_clean'], r['prod'] or '',
                   r['amt_f'], r['note'] or '', r['src']])
    for i, w in enumerate([6, 18, 8, 26, 40, 12, 16, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(cols)):
        for c in row:
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            c.alignment = Alignment(vertical='center')
        row[5].number_format = '#,##0.00;[Red]-#,##0.00'
        if row[2].value == '退款':
            row[2].font = Font(color='C00000')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    ws2 = wb.create_sheet('商户汇总')
    ws2.append(['商户名称', '笔数', '合计金额(元)'])
    agg = defaultdict(lambda: [0, 0.0])
    for r in all_rows:
        a = agg[r['shop_clean']]
        a[0] += 1
        a[1] += r['amt_f'] or 0
    for k, v in sorted(agg.items(), key=lambda x: x[1][1]):
        ws2.append([k, v[0], round(v[1], 2)])
    ws2.append(['总计', len(all_rows), round(sum(r['amt_f'] or 0 for r in all_rows), 2)])
    for i, w in enumerate([32, 10, 18], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for c in ws2[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT

    ws3 = wb.create_sheet('月度汇总')
    ws3.append(['月份', '笔数', '支出合计(元)', '退款笔数'])
    mon = defaultdict(lambda: [0, 0.0, 0])
    for r in all_rows:
        k = (r['dt'] or '未知')[:7]
        m = mon[k]
        m[0] += 1
        m[1] += (r['amt_f'] or 0)
        if r['type'] == '退款':
            m[2] += 1
    for k in sorted(mon, reverse=True):
        v = mon[k]
        ws3.append([k, v[0], round(v[1], 2), v[2]])
    for i, w in enumerate([14, 10, 18, 12], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for c in ws3[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT

    ws4 = wb.create_sheet('处理报告')
    ws4.append(['文件名', '状态', '说明'])
    for f, st, msg in bad_files:
        ws4.append([f, st, msg])
    for i, w in enumerate([40, 12, 60], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    for c in ws4[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT

    wb.save(outpath)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('folder', help='存放图片的文件夹')
    ap.add_argument('--out', default=None, help='输出 xlsx 路径')
    ap.add_argument('--recursive', action='store_true', help='同时扫描子文件夹')
    args = ap.parse_args()

    folder = os.path.abspath(args.folder)
    files = []
    if args.recursive:
        for root, _, names in os.walk(folder):
            files += [os.path.join(root, n) for n in names if n.lower().endswith(IMG_EXT)]
    else:
        files = [os.path.join(folder, n) for n in os.listdir(folder)
                 if n.lower().endswith(IMG_EXT)]
    files.sort()
    if not files:
        print('未找到图片文件')
        return
    print(f'发现 {len(files)} 张图片\n' + '=' * 60)

    engine = RapidOCR()
    all_rows, bad_files, per_image = [], [], []

    def key(r):
        return (r['dt'], r['shop_clean'], r['amt_f'])

    # 第一阶段：逐张解析
    for i, f in enumerate(files, 1):
        name = os.path.basename(f)
        try:
            ok, msg, im = check_readable(f)
            if not ok:
                bad_files.append((name, '跳过', msg))
                print(f'[{i}/{len(files)}] {name} -> 跳过：{msg}')
                continue
            boxes = ocr_image(im, engine)
            lines = dedupe_boxes(boxes)
            left, right = detect_columns(lines, im.size[0])
            rows = [r for r in parse_records(lines, left, right, name)
                    if r['dt'] and r['amt_f'] is not None]
            per_image.append((name, rows))
            print(f'[{i}/{len(files)}] {name} -> {msg}；识别 {len(rows)} 条')
        except Exception as e:
            bad_files.append((name, '失败', f'{type(e).__name__}: {e}'))
            print(f'[{i}/{len(files)}] {name} -> 失败：{e}')
            traceback.print_exc()

    # 第二阶段：跨图去重（同一分钟同店同额的重复下单要保留，故按出现次数取最大值）
    mult = defaultdict(int)
    for _, rows in per_image:
        cnt = Counter(key(r) for r in rows)
        for k, v in cnt.items():
            mult[k] = max(mult[k], v)
    emitted = Counter()
    for _, rows in per_image:
        for r in rows:
            k = key(r)
            if emitted[k] < mult[k]:
                emitted[k] += 1
                all_rows.append(r)
    deduped = sum(len(rows) for _, rows in per_image) - len(all_rows)
    print(f'跨图去重：移除 {deduped} 条重复（重叠截图）')

    all_rows.sort(key=lambda r: r['dt'], reverse=True)
    for i, r in enumerate(all_rows, 1):
        r['seq'] = i

    out = args.out or os.path.join(folder, '账单汇总.xlsx')
    export(all_rows, bad_files, out)
    total = sum(r['amt_f'] for r in all_rows)
    print('=' * 60)
    print(f'完成：{len(all_rows)} 条记录，合计 ¥{abs(total):,.2f}')
    if all_rows:
        print(f'时间范围：{all_rows[-1]["dt"]} ~ {all_rows[0]["dt"]}')
    print(f'跳过/失败 {len(bad_files)} 个文件')
    print(f'输出：{out}')


if __name__ == '__main__':
    main()
